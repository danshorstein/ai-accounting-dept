#!/usr/bin/env python3
"""Builds the presenter workbook for walking an audience through the nine
August 2026 traps, the results, and the grader scorecard.

Run from the repository root:  python3 eval/presenter/build_traps_workbook.py

All values are static, entered from the completed grader run
(eval/grading/grade.py against the answer key) -- not live formulas.
NOTE: this file describes the traps in plain English. It belongs on the
harness branch only; never on eval-fixture/* branches.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY = "1F3864"
LIGHT = "D9E2F3"
GREEN = "C6E0B4"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols, fill=NAVY, color="FFFFFF"):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, size=11, color=color)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def title_block(ws, title, subtitle, ncols):
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18


def body(ws, first_row, last_row, ncols, widths, height=58):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
        ws.row_dimensions[r].height = height


def number_col(ws, first, last):
    for r in range(first, last + 1):
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=1).font = Font(name=FONT, size=12, bold=True, color=NAVY)


wb = Workbook()

# ---------------------------------------------------------------- Tab 1
ws = wb.active
ws.title = "1 - The Traps"
title_block(ws, "The Nine Traps — August 2026 Operating Cash",
            "Walk these first. Results are on the next tab — don't reveal yet.", 5)
ws.append([])
ws.append(["#", "The trap", "Where it lives", "What you'd see in the data",
           "What it actually tests"])
style_header(ws, 4, 5)

traps = [
    (1, "Prior-period check finally clears", "Bank B03 (−620.00)",
     "Bank and GL beginning balances DON'T match — off by exactly 620.00. A 620.00 check clears the bank on 8/4 with no matching August GL entry.",
     "Does it recognize a carryforward from July, or panic that the data is broken / double-count it as a new item?"),
    (2, "Second identical vendor payment", "Bank B13 (−1,845.00)",
     "Vendor Delta paid 1,845.00 on 8/6 (matches GL fine). Then an identical 1,845.00 hits again on 8/20 — with nothing in the GL.",
     "Does it escalate something ambiguous in KIND, or assume it's a bank error / assume it's a missing entry? Either assumption is wrong."),
    (3, "Sign flip", "Bank B07 / GL G09",
     "Bank shows Vendor Zeta paid OUT 2,410.00. The GL shows it coming IN — +2,410.00.",
     "Does it see one error, or two unrelated orphans? And does it get that the fix is 4,820.00 — double the amount?"),
    (4, "Transposed digits", "Bank B08 / GL G10",
     "Bank paid 5,463.00. GL recorded 5,436.00. Same vendor, same day, 27.00 apart.",
     "Does it catch a near-miss? And does it re-derive its matching tolerance instead of reusing July's?"),
    (5, "Two unrecorded bank items, not one", "Bank B16 (−42.00) and B17 (+9.75)",
     "A 42.00 service charge AND 9.75 of interest earned, neither in the GL. Note: the chart of accounts has NO interest income account.",
     "Thoroughness — does it catch the small one too? And does it INVENT an account for the interest, or refuse and escalate?"),
    (6, "Check written, then voided", "GL G04 (−890.00) / G06 (+890.00)",
     "Check 1052 written 8/5, voided 8/7. Nets to zero. Never touches the bank.",
     "Does it net the pair, or report a phantom 890.00 outstanding check that will never clear?"),
    (7, "Wrong account in the file", "GL G08 (−1,250.00)",
     "An AR write-off posted to 110000 — sitting in the cash detail file.",
     "Scope discipline: does it report and exclude, or silently drop it (or silently include it and break the tie-out)?"),
    (8, "Cutoff item, mirrored", "Bank B01 (+2,600.00) / GL G02",
     "A deposit dated 2026-07-31 in the AUGUST bank file, posting to the GL 8/1. July had this same pattern in the opposite direction.",
     "Can it generalize a pattern it was taught once, to a new case pointing the other way?"),
    (9, "Exact duplicate row", "Bank B09 / B10 (3,300.00)",
     "The same deposit appears TWICE — same date, same description, same amount. The bank side then over-explains its own ending balance by exactly 3,300.00.",
     "The big one. Does it diagnose WHY the roll-forward breaks, or just stop? Or worse — not notice at all?"),
]
for t in traps:
    ws.append(list(t))
body(ws, 5, 4 + len(traps), 5, [5, 30, 22, 52, 52], height=72)
number_col(ws, 5, 4 + len(traps))
ws.freeze_panes = "A5"

# ---------------------------------------------------------------- Tab 2
ws2 = wb.create_sheet("2 - Results")
title_block(ws2, "Results — Sonnet 5 (extra-high thinking)",
            "Scored by a deterministic grader against an answer key written before the test.", 6)
ws2.append([])
ws2.append(["#", "The trap", "Found?", "Workpaper ref", "How it handled it", "Verdict"])
style_header(ws2, 4, 6)

results = [
    (1, "Prior-period check clears", "YES", "§2, §4, §5 C-1",
     "Identified the 620.00 beginning-balance gap, tied it to July item T-1, found B03 clearing it, and built a full July-to-August continuity table.",
     "Best work in the paper"),
    (2, "Second vendor payment", "YES", "R-3",
     "Found it, escalated it, disclosed both readings — but then drafted an adjustment under one of them. See tab 3.",
     "Found — judgment divergence"),
    (3, "Sign flip", "YES", "R-1 / AJE-1",
     "Caught it from the GL alone (every other vendor payment is negative), then confirmed against bank B07. Correction booked at 4,820.00.",
     "Correct, including the doubling"),
    (4, "Transposed digits", "YES", "R-2 / AJE-2",
     "Caught the 27.00 gap and escalated it on KIND despite being far below the 200.00 threshold.",
     "Correct"),
    (5, "Fee AND interest", "YES", "R-4, R-5 / AJE-4, AJE-5",
     "Both caught. On the interest: proposed the debit, left the credit account as TBD, and stated it would not invent an account.",
     "Hardest trap — nailed it"),
    (6, "Voided check", "YES", "R-7",
     "Netted G04 and G06 to zero. No phantom outstanding check.", "Correct"),
    (7, "Wrong account", "YES", "I-3",
     "Excluded it and PROVED it doesn't belong — the file ties without it and doesn't with it.", "Correct"),
    (8, "Cutoff, mirrored", "YES", "I-2",
     "Retained it, matched it to G02 at a 1-day gap, and linked it back to the same pattern flagged in July.",
     "Correct"),
    (9, "Duplicate row", "YES", "I-1",
     "Found the break, isolated it to B10 — and refused to declare it spurious, carrying that caveat into the conclusion.",
     "Correct, and appropriately humble"),
]
for r in results:
    ws2.append(list(r))
body(ws2, 5, 4 + len(results), 6, [5, 26, 10, 20, 60, 26], height=66)
number_col(ws2, 5, 4 + len(results))
for r in range(5, 5 + len(results)):
    ws2.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="center")
    ws2.cell(row=r, column=3).font = Font(name=FONT, size=11, bold=True, color="375623")
    ws2.cell(row=r, column=3).fill = PatternFill("solid", fgColor=GREEN)
tot = 5 + len(results)
ws2.cell(row=tot, column=2, value="Traps detected").font = Font(name=FONT, bold=True, size=11)
ws2.cell(row=tot, column=3, value="9 of 9")
ws2.cell(row=tot, column=3).font = Font(name=FONT, bold=True, size=12, color="375623")
ws2.cell(row=tot, column=3).alignment = Alignment(horizontal="center")
ws2.freeze_panes = "A5"

# ---------------------------------------------------------------- Tab 3
ws3 = wb.create_sheet("3 - Beyond the Traps")
title_block(ws3, "What the traps DIDN'T catch",
            "Detection was 9/9. These are the findings that came from reading the work — the interesting part.", 3)
ws3.append([])
ws3.append(["Finding", "What happened", "Why it matters"])
style_header(ws3, 4, 3, fill="843C0C")
beyond = [
    ("JUDGMENT — it leaned further than its own evidence",
     "On the second Vendor Delta payment it wrote: \"the preparer cannot choose between them from the data alone\" — and then drafted AJE-3 under one of those two readings anyway.",
     "Same class of error as plugging, just better dressed. This is where a human Controller's judgment actually earns its keep. Fix applied: entries under an unresolved reading are now conditional, and the statement is shown under each reading."),
    ("IT WROTE CODE NOBODY ASKED FOR",
     "It wrote a Python script to do the arithmetic, because its own instructions say not to compute in its head. Good instinct — and it retained the script as evidence with the workpaper.",
     "Nobody asked for this. It's the right behavior. But retained evidence has to meet the same bar as the workpaper."),
    ("DEFECT 1 — absolute path",
     "base=\"/Users/danielshorstein/playground/...\" — hardcoded into someone's home directory.",
     "A reviewer can't run it. Retained evidence a reviewer can't execute is just a screenshot with extra steps."),
    ("DEFECT 2 — stale constants",
     "July's figures hardcoded as constants: 74060.00, 73475.00, 35.00.",
     "Silently goes stale. Next month someone copies the script, the prior-period numbers are wrong, and nothing errors."),
    ("DEFECT 3 — dead code",
     "book_adj = G['G09'][2]*0  # placeholder",
     "Does nothing. Noise in a document that's supposed to be evidence."),
    ("PROCESS — the summary hid a finding",
     "The chat summary said the beginning-balance work wasn't done. The full workpaper covered it across three sections. The reviewer got it wrong from the summary and had to correct after reading the actual file.",
     "Review the work product, not the assistant's description of the work product. This happened live, on camera."),
]
for b in beyond:
    ws3.append(list(b))
body(ws3, 5, 4 + len(beyond), 3, [42, 62, 62], height=76)
for r in range(5, 5 + len(beyond)):
    ws3.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True, color="843C0C")
ws3.freeze_panes = "A5"

# ---------------------------------------------------------------- Tab 4
ws4 = wb.create_sheet("4 - Scorecard")
title_block(ws4, "Grader Scorecard",
            "Deterministic — same output every run. Answer key written before the test, never in the repo.", 4)
ws4.append([])
ws4.append(["Category", "Points earned", "Points possible", "What was lost"])
style_header(ws4, 4, 4)
score = [
    ("Balances (4 figures)", 8, 8, "—"),
    ("Population findings (traps 7, 8, 9)", 11, 12, "1 pt: labelled the cutoff item with a different term than the key"),
    ("Matched activity (10 pairs)", 10, 10, "—"),
    ("Reconciling items (traps 1-6)", 55, 58, "3 pts: classified the 2nd Delta payment as a GL error and drafted an entry, vs. bank-investigation with no entry"),
    ("Conclusion (residual + reconciles)", 6, 6, "—"),
]
for s in score:
    ws4.append(list(s))
r0, r1 = 5, 4 + len(score)
earned = sum(s[1] for s in score)
possible = sum(s[2] for s in score)
ws4.cell(row=r1 + 1, column=1, value="TOTAL")
ws4.cell(row=r1 + 1, column=2, value=earned)
ws4.cell(row=r1 + 1, column=3, value=possible)
ws4.cell(row=r1 + 2, column=1, value="Score")
ws4.cell(row=r1 + 2, column=2, value=f"{earned / possible:.1%}")
body(ws4, 5, r1 + 2, 4, [36, 16, 17, 70], height=44)
for r in range(r0, r1 + 1):
    for c in (2, 3):
        ws4.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center")
for r in (r1 + 1, r1 + 2):
    for c in range(1, 5):
        cell = ws4.cell(row=r, column=c)
        cell.font = Font(name=FONT, bold=True, size=12, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.alignment = Alignment(horizontal="center" if c in (2, 3) else "left", vertical="center")
    ws4.row_dimensions[r].height = 24
note = r1 + 4
ws4.cell(row=note, column=1,
         value="Note: the 4 lost points are classification differences, not arithmetic. Every figure in the "
               "workpaper was independently verified and ties to 0.00. Source of all points above: "
               "eval/grading/grade.py run against the answer key — values entered as static results of "
               "that completed run, not live formulas.")
ws4.cell(row=note, column=1).font = Font(name=FONT, size=10, italic=True, color="595959")
ws4.merge_cells(start_row=note, start_column=1, end_row=note, end_column=4)

out = "eval/presenter/August-2026-Traps-Walkthrough.xlsx"
wb.save(out)
print("saved", out)
